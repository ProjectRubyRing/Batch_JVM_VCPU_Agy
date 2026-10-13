import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def build_workbook():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default sheet

    if wb._fonts:
        wb._fonts[0].name = "Meiryo UI"

    FONT_NAME = "Meiryo UI"

    # Monotone Palette Definitions (Grayscale / Charcoal / Slate)
    fill_charcoal_dark = PatternFill(start_color="212529", end_color="212529", fill_type="solid")  # Primary Header (#212529)
    fill_slate_med = PatternFill(start_color="495057", end_color="495057", fill_type="solid")       # Section Header (#495057)
    fill_silver_sub = PatternFill(start_color="6C757D", end_color="6C757D", fill_type="solid")      # Sub Header (#6C757D)
    
    fill_light_gray = PatternFill(start_color="E9ECEF", end_color="E9ECEF", fill_type="solid")      # Category / Highlight (#E9ECEF)
    fill_soft_gray = PatternFill(start_color="F1F3F5", end_color="F1F3F5", fill_type="solid")       # Soft Accent (#F1F3F5)
    fill_zebra_even = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")      # Table Even Row (#F8F9FA)
    fill_zebra_odd = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")       # Table Odd Row (#FFFFFF)
    fill_calc_total = PatternFill(start_color="DEE2E6", end_color="DEE2E6", fill_type="solid")      # Total / KPI (#DEE2E6)

    # Fonts (STRICTLY Meiryo UI across all variants)
    font_title = Font(name=FONT_NAME, size=15, bold=True, color="111111")
    font_subtitle = Font(name=FONT_NAME, size=9.5, italic=True, color="555555")
    font_sec_hdr = Font(name=FONT_NAME, size=11, bold=True, color="FFFFFF")
    font_tbl_hdr = Font(name=FONT_NAME, size=9.5, bold=True, color="FFFFFF")
    font_body = Font(name=FONT_NAME, size=9, color="212529")
    font_body_bold = Font(name=FONT_NAME, size=9, bold=True, color="111111")
    font_body_code = Font(name=FONT_NAME, size=9, bold=True, color="2B303A")
    font_note = Font(name=FONT_NAME, size=8.5, italic=True, color="6C757D")

    # Borders
    thin_border_side = Side(style="thin", color="CED4DA")
    thick_bottom_side = Side(style="medium", color="212529")
    double_bottom_side = Side(style="double", color="212529")

    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    border_tbl_hdr = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thick_bottom_side)
    border_total = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=double_bottom_side)

    # Alignments
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)

    def style_header_row(ws, row_idx, headers, fill=fill_charcoal_dark, font=font_tbl_hdr, border=border_tbl_hdr):
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=h)
            cell.fill = fill
            cell.font = font
            cell.alignment = align_center
            cell.border = border
        ws.row_dimensions[row_idx].height = 26

    def enforce_meiryo_all(ws):
        for row in ws.iter_rows():
            for cell in row:
                cur = cell.font
                if cur:
                    cell.font = Font(
                        name=FONT_NAME,
                        size=cur.size,
                        bold=cur.bold,
                        italic=cur.italic,
                        color=cur.color
                    )
                else:
                    cell.font = Font(name=FONT_NAME, size=9, color="212529")

    def auto_fit_columns(ws, max_cols=10, min_width=12, max_width=68):
        for col in range(1, max_cols + 1):
            col_letter = get_column_letter(col)
            max_len = 0
            for row in range(1, ws.max_row + 1):
                val = ws.cell(row=row, column=col).value
                if val:
                    val_str = str(val)
                    line_lens = [sum(2 if ord(c) > 127 else 1 for c in line) for line in val_str.split("\n")]
                    max_len = max(max_len, max(line_lens) if line_lens else 0)
            target_width = min(max(max_len + 3, min_width), max_width)
            ws.column_dimensions[col_letter].width = target_width

    # =============================================================
    # Sheet 1: 表紙・全体概要 (Cover & Executive Summary)
    # =============================================================
    ws1 = wb.create_sheet(title="表紙・全体概要")
    ws1.views.sheetView[0].showGridLines = True

    ws1.merge_cells("A1:G1")
    ws1["A1"] = "EC2 (RHEL 9.8) × OpenJDK 21 親子JVMバッチ並列処理 性能・リソース設計書"
    ws1["A1"].font = font_title
    ws1["A1"].alignment = align_left
    ws1.row_dimensions[1].height = 32

    ws1.merge_cells("A2:G2")
    ws1["A2"] = "～ プロセス分離、メモリ構造、vCPUスケジューリング、sar調査コマンド、サイジング計算完全版 ～"
    ws1["A2"].font = font_subtitle
    ws1["A2"].alignment = align_left
    ws1.row_dimensions[2].height = 20

    meta = [
        ("対象環境 インフラ", "AWS EC2 (Nitro System / Intel Xeon or AMD EPYC / gp3 EBS)"),
        ("対象環境 OS", "Red Hat Enterprise Linux 9.8 (Kernel 5.14+ / cgroups v2 / systemd 252+)"),
        ("Java ランタイム", "OpenJDK 21 (LTS) 64-Bit Server VM (G1GC標準 / Generational ZGC対応)"),
        ("バッチ処理方式", "親JVM (1プロセス：司令塔・管理) ＋ 子JVM (10並列プロセス：業務実処理)"),
        ("総プロセス構成", "Javaプロセス 計11個 (親1 ＋ 子10) ＋ OSデーモン・カーネルスレッド群"),
        ("設計ポリシー", "OSプロセス分離による完全な障害隔離(Fault Isolation)と高スループットの両立"),
        ("フォント・配色規定", "全セル統一フォント: Meiryo UI ／ カラーパレット: モノトーン (Charcoal & Grayscale)")
    ]

    for r_idx, (k, v) in enumerate(meta, 4):
        ws1.cell(row=r_idx, column=1, value=k).font = font_body_bold
        ws1.cell(row=r_idx, column=1).fill = fill_light_gray
        ws1.cell(row=r_idx, column=1).border = border_cell
        ws1.cell(row=r_idx, column=1).alignment = align_left
        ws1.merge_cells(start_row=r_idx, start_column=2, end_row=r_idx, end_column=7)
        c2 = ws1.cell(row=r_idx, column=2, value=v)
        c2.font = font_body
        c2.alignment = align_left
        for c in range(2, 8):
            ws1.cell(row=r_idx, column=c).border = border_cell
        ws1.row_dimensions[r_idx].height = 21

    idx_headers = ["No.", "シート名", "主な解説内容・提供価値", "想定利用者 / 活用シーン"]
    style_header_row(ws1, 12, idx_headers, fill=fill_charcoal_dark)
    ws1.merge_cells("D12:G12")

    sheets_index = [
        ("1", "表紙・全体概要", "システム前提条件、基本構成、本設計書の目次と核心サマリ", "ステークホルダー、開発・運用チーム全体"),
        ("2", "小学生でもわかる動作イメージ", "「巨大な学校と校長先生、10の特別教室」に例えた直感的メンタルモデル", "新人教育、非IT関係者への説明、概念把握"),
        ("3", "全体アーキテクチャ", "親JVM(1)と子JVM(10)の連携詳細、OS/Nitroとの階層モデルと障害波及境界", "インフラ設計者、システムアーキテクト"),
        ("4", "プロセスメモリ・vCPU内部構造", "JVM内部メモリ(Heap/Non-Heap)詳細内訳、OS Native Thread、CFSスケジューリング", "JVM性能エンジニア、基盤設計担当者"),
        ("5", "性能試験コマンド完全詳解", "sar(CPU, Load, Memory, Swap, Paging, Switch, Disk)・pidstat・vmstat詳解", "性能試験担当者、SRE、性能ボトルネック分析"),
        ("6", "サイジング計算シミュレータ", "10並列時の実メモリ(RSS)・Page Cache・vCPU所要量試算シート(計算式組込)", "インフラ調達、EC2インスタンス選定、キャパシティ計画"),
        ("7", "歴史的背景・技術選定理由", "単一JVMスレッド型の挫折(OOM/GC凍結/リーク)とマルチプロセス分離の必然性", "方式設計書レビュー、アーキテクチャ妥当性確認"),
        ("8", "チューニング・推奨設定パラメータ", "OpenJDK 21推奨JVMオプション、OSカーネルパラメータ(sysctl)、CDS高速化", "本番環境構築、性能最適化チューニング")
    ]

    for i, row in enumerate(sheets_index, 13):
        fill_cur = fill_zebra_even if i % 2 == 0 else fill_zebra_odd
        ws1.cell(row=i, column=1, value=row[0]).alignment = align_center
        ws1.cell(row=i, column=2, value=row[1]).alignment = align_left
        ws1.cell(row=i, column=3, value=row[2]).alignment = align_left
        ws1.merge_cells(start_row=i, start_column=4, end_row=i, end_column=7)
        ws1.cell(row=i, column=4, value=row[3]).alignment = align_left
        
        for c in range(1, 8):
            cell = ws1.cell(row=i, column=c)
            cell.font = font_body_bold if c in (1, 2) else font_body
            cell.fill = fill_cur
            cell.border = border_cell
        ws1.row_dimensions[i].height = 22

    ws1.merge_cells("A22:G22")
    ws1["A22"] = "【本アーキテクチャの最重要設計ポイント（サマリ）】"
    ws1["A22"].font = font_sec_hdr
    ws1["A22"].fill = fill_slate_med
    ws1["A22"].alignment = align_left
    ws1.row_dimensions[22].height = 24

    sum_desc = (
        "◆ 完全な耐障害性の獲得: 親JVM(管理)と子JVM(業務)をプロセス分離することで、特定バッチのOutOfMemoryErrorやSIGSEGVクラッシュが"
        "他9本の子JVMや親JVMに波及せず、親JVMが exit code (137等) を検知して自律的にリトライ・隔離制御を行える。\n"
        "◆ 実メモリ(RSS)サイジングの厳格化: 「子JVM 10並列のメモリ ＝ 10 × -Xmx」ではない。Metaspace、CodeCache、スレッドスタック(-Xss)、"
        "GC内部構造体、glibc mallocによるNativeメモリが10プロセス分重複するため、1子JVMあたりHeap+約0.8～1.2GBのRSSを見込む必要がある。\n"
        "◆ vCPUとスレッド競合の制御: 11個のJVMが各々デフォルト数のGCスレッド(ParallelGCThreads)やJITスレッドを起動すると、総スレッド数が数百本に跳ね上がり"
        "CPUコンテキストスイッチ(sar -w の cswch/s)が爆発する。子JVM起動引数で並列スレッド数を制限することが性能試験成功の鍵となる。"
    )
    ws1.merge_cells("A23:G23")
    ws1["A23"] = sum_desc
    ws1["A23"].font = font_body
    ws1["A23"].alignment = align_left
    ws1.row_dimensions[23].height = 68
    for c in range(1, 8):
        ws1.cell(row=23, column=c).border = border_cell

    auto_fit_columns(ws1, max_cols=7)
    ws1.column_dimensions["A"].width = 8
    ws1.column_dimensions["B"].width = 25
    ws1.column_dimensions["C"].width = 44
    ws1.column_dimensions["D"].width = 30

    # =============================================================
    # Sheet 2: 小学生でもわかる動作イメージ (Analogy & Mental Model)
    # =============================================================
    ws2 = wb.create_sheet(title="小学生でもわかる動作イメージ")
    ws2.views.sheetView[0].showGridLines = True

    ws2.merge_cells("A1:G1")
    ws2["A1"] = "【たとえ話で直感理解】巨大な学校と校長先生、そして10の特別教室"
    ws2["A1"].font = font_title
    ws2.row_dimensions[1].height = 30

    ws2.merge_cells("A2:G2")
    ws2["A2"] = "難しい専門用語を「学校の校舎・先生・生徒・机・警備員」に置き換えて、直感的なメンタルモデルで完全理解します。"
    ws2["A2"].font = font_subtitle
    ws2.row_dimensions[2].height = 18

    headers2 = ["IT専門用語", "学校にたとえると何？", "具体的なお仕事・役割", "トラブルが起きたときの挙動", "なぜこの仕組みが必要なの？（メリット）"]
    style_header_row(ws2, 4, headers2, fill=fill_charcoal_dark)
    ws2.merge_cells("E4:G4")

    analogy_items = [
        ("EC2 インスタンス\n(仮想サーバー)", "大きな学校の校舎全体", "生徒たちが作業するための建物全体。建物の広さ（メモリ）と、働ける先生の頭脳の数（vCPU）があらかじめ決まっている。", "校舎ごと停電すると、学校全体の全作業が一瞬で止まります。", "クラウド上に自分たち専用の頑丈で安全な校舎を借りて作業できる。"),
        ("vCPU\n(仮想CPUコア)", "教室で働く先生の「頭脳と両手」", "計算ドリルを解いたり、文字を書いたり、実際に頭と手を動かして作業を進める力。", "作業が多すぎて頭がパンクすると、生徒たちの順番待ち（CPU待ち行列）が長くなります。", "頭や手の数（vCPU数）が多ければ多いほど、たくさんの仕事を同時に進められます。"),
        ("物理メモリ (RAM)", "教室にある「大きな作業机」", "今すぐ使うノートや筆記用具を広げておく場所。机が広いほどたくさんの資料を一度に広げられる。", "机の上が物で溢れかえると、作業ができなくなって警備員に教室から追い出されます（OOM Killer）。", "引き出し（SSD/ハードディスク）から取り出すより、机の上（メモリ）にある方が10万倍速く作業できる！"),
        ("スワップ領域\n(Swap)", "廊下に置かれた「荷物の一時置きダンボール」", "机がいっぱいになった時、たまにしか使わないノートを廊下の箱に一時退避させる仕組み。", "廊下までノートを取りに行くとすごく時間がかかるので、作業スピードがカメのように遅くなります。", "完全に机から溢れて作業が全滅するのを、スピードを極端に落としてでもギリギリ防ぐ安全装置。"),
        ("親JVM\n(Master Process)", "職員室にいる「校長先生」", "「次はこのバッチ宿題をやってね」と指示を出し、10個の特別教室を見守る司令塔。自分では絵の具を使いません。", "校長先生が倒れると次の宿題の指示が出せなくなります（でも自分は重労働をしないので滅多に倒れません）。", "校長先生自身は絶対に汚れる作業をしないことで、学校全体の指揮が絶対にストップしないように守る。"),
        ("子JVM (10個並列)\n(Worker Processes)", "10個の「独立した特別工作教室」", "校長先生から宿題をもらって、実際に絵の具やノコギリを使って重い工作（バッチ実処理）を集中して行う部屋。", "1つの教室で絵の具をぶちまけて大爆発（メモリパンク）しても、その部屋だけが片付けられます。", "【最大のポイント！】他の9個の教室と職員室は頑丈な壁で仕切られているので、汚れが絶対に飛び散らない！"),
        ("JVMヒープメモリ\n(-Xmx)", "各教室の中に置かれた「生徒の個人机」", "生徒が絵の具や画用紙を広げるスペース。この広さを超えて画用紙を広げようとすると怒られます。", "机の端を超えて物を置こうとすると「もう置けません！（OutOfMemoryError）」と叫んで止まります。", "生徒ごとに「あなたの机はこの広さまでね」と上限を決めておくことで、校舎全体の机を独り占めさせない。"),
        ("JVMメタスペース\n(Metaspace)", "教室の壁にある「教科書の本棚」", "プログラムのルールやクラスの設計図（教科書）を置いておく場所。生徒の個人机とは別の専用の本棚。", "本棚が溢れると、新しい教科書が読めなくなって作業が止まります。", "生徒の作業机（ヒープ）と、教科書の棚（メタスペース）を分けることで整理整頓される。"),
        ("OS (Linuxカーネル)", "学校の「厳格な警備員兼管理人」", "誰がどの頭脳（CPU）を何秒使うか秒単位で交通整理し、机の空き具合を24時間見回る管理人。", "誰かが暴走して学校のルールを壊しそうになると、強制的に部屋から叩き出します（kill -9）。", "みんなが勝手なことをして学校全体が崩壊しないよう、公平かつ厳しく秩序を守る。"),
        ("Linux OOM Killer", "暴走した生徒を強制退去させる「警備員」", "学校全体の机（メモリ）が本当にゼロになった時、一番机を散らかしている部屋を強制閉鎖する。", "警告なしで一瞬で部屋から叩き出されます（exit code 137）。", "学校全体が倒壊（OSクラッシュ）するのを防ぐため、一番机を散らかした部屋を身代わりに倒す苦肉の策。"),
        ("sar コマンド", "学校全体の「監視カメラと健康診断計」", "「先生の頭は疲れてないか？」「机はどれくらい埋まってるか？」「廊下を走ってる人は何人か？」を1秒ごとに記録。", "過去の記録をいつでも見返して、「昨日の夜22時に誰がサボっていたか・苦しんでいたか」を特定できます。", "どこが苦しいのか（CPUなのか、メモリなのか、机の広さなのか）が一目でわかる！")
    ]

    for r_idx, r_val in enumerate(analogy_items, 5):
        fill_cur = fill_light_gray if "子JVM" in r_val[0] or "親JVM" in r_val[0] else (fill_zebra_even if r_idx % 2 == 0 else fill_zebra_odd)
        for c_idx, val in enumerate(r_val[:4], 1):
            cell = ws2.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_body_bold if c_idx in (1, 2) else font_body
            cell.fill = fill_cur
            cell.border = border_cell
            cell.alignment = align_center if c_idx in (1, 2) else align_left
        
        ws2.merge_cells(start_row=r_idx, start_column=5, end_row=r_idx, end_column=7)
        c5 = ws2.cell(row=r_idx, column=5, value=r_val[4])
        c5.font = font_body
        c5.alignment = align_left
        for c in range(5, 8):
            ws2.cell(row=r_idx, column=c).border = border_cell
            ws2.cell(row=r_idx, column=c).fill = fill_cur
        ws2.row_dimensions[r_idx].height = 36

    ws2.merge_cells("A17:G17")
    ws2["A17"] = "【核心の疑問】なぜ「1つの大部屋で10人一緒に作業」ではなく「10の小部屋に分ける」のか？"
    ws2["A17"].font = font_sec_hdr
    ws2["A17"].fill = fill_slate_med
    ws2["A17"].alignment = align_left
    ws2.row_dimensions[17].height = 24

    ws2.merge_cells("A18:C18")
    ws2["A18"] = "方式A: 1つの大部屋で10人作業 (単一JVMスレッド並列)"
    ws2["A18"].font = font_tbl_hdr
    ws2["A18"].fill = fill_charcoal_dark
    ws2["A18"].alignment = align_center

    ws2.merge_cells("D18:G18")
    ws2["D18"] = "方式B: 10の独立した小部屋で作業 (本システムの親子JVM方式)"
    ws2["D18"].font = font_tbl_hdr
    ws2["D18"].fill = fill_charcoal_dark
    ws2["D18"].alignment = align_center
    ws2.row_dimensions[18].height = 24

    comparisons = [
        ("メリット", "・部屋が1つだけなので、準備や片付けが一瞬（起動が速い）\n・教科書棚（メタスペース等）が1つで済むのでメモリ消費が少ない",
         "・【圧倒的な安全性】1人が机を汚しても、他の9部屋は1ミリも汚れない\n・作業が終わったら部屋ごと丸ごと取り壊して新品にするのでゴミが100%残らない"),
        ("デメリット", "・【致命的】1人がインクを床にこぼすと、全員の足元が汚れて全員の作業が中断する\n・1人が机を散らかしすぎると部屋全体が閉鎖され全員退場になる（連座倒れ）",
         "・部屋を10個建てるので、それぞれの部屋に教科書棚（メタスペース等）が必要で少しメモリを多く喰う\n・毎回部屋を新しく建てるので、少しだけ準備時間（起動時間）がかかる"),
        ("小学生向け結論", "「仲良く1つの部屋でやると、1人の失敗でみんな連帯責任で怒られる！」",
         "「壁で仕切られた別々の部屋でやれば、隣の人が失敗しても自分は安心して最後まで宿題を終わらせられる！」")
    ]

    for c_i, comp in enumerate(comparisons, 19):
        ws2.merge_cells(start_row=c_i, start_column=1, end_row=c_i, end_column=3)
        ws2.cell(row=c_i, column=1, value=f"{comp[0]}: {comp[1]}").font = font_body
        ws2.cell(row=c_i, column=1).alignment = align_left
        for c in range(1, 4):
            ws2.cell(row=c_i, column=c).border = border_cell
            ws2.cell(row=c_i, column=c).fill = fill_zebra_even if c_i % 2 == 0 else fill_zebra_odd

        ws2.merge_cells(start_row=c_i, start_column=4, end_row=c_i, end_column=7)
        ws2.cell(row=c_i, column=4, value=f"{comp[0]}: {comp[2]}").font = font_body_bold if "安全性" in comp[2] or "結論" in comp[0] else font_body
        ws2.cell(row=c_i, column=4).alignment = align_left
        for c in range(4, 8):
            ws2.cell(row=c_i, column=c).border = border_cell
            ws2.cell(row=c_i, column=c).fill = fill_calc_total if "結論" in comp[0] else (fill_zebra_even if c_i % 2 == 0 else fill_zebra_odd)
        ws2.row_dimensions[c_i].height = 42

    auto_fit_columns(ws2, max_cols=7)
    ws2.column_dimensions["A"].width = 22
    ws2.column_dimensions["B"].width = 25
    ws2.column_dimensions["C"].width = 38
    ws2.column_dimensions["D"].width = 30

    # =============================================================
    # Sheet 3: 全体アーキテクチャ (Architecture)
    # =============================================================
    ws3 = wb.create_sheet(title="全体アーキテクチャ")
    ws3.views.sheetView[0].showGridLines = True

    ws3.merge_cells("A1:H1")
    ws3["A1"] = "システム階層構造とプロセス連携メカニズム (1親JVM ＋ 10子JVM並列)"
    ws3["A1"].font = font_title
    ws3.row_dimensions[1].height = 30

    headers3_arch = ["レイヤー", "構成コンポーネント", "プロセス種別 / PID例", "主な責務・動作内容", "割当リソース", "プロセス間通信 / 連携手段", "障害時の影響範囲", "備考 / 最新動向"]
    style_header_row(ws3, 3, headers3_arch, fill=fill_charcoal_dark)

    arch_details = [
        ("ハードウェア・仮想化", "AWS EC2 (Nitro System)", "Nitro Card / Hypervisor", "物理CPU・メモリの仮想化、EBS/VPCハードウェア完全オフロード", "物理コア, 物理RAM, NVMe EBS", "PCIeエミュレーション, SR-IOV", "インスタンス全体停止", "Nitroによりジッター極小化・ベアメタル同等の安定性"),
        ("OSカーネル", "RHEL 9.8 (Linux Kernel 5.14+)", "カーネル空間 (PID 0)", "プロセススケジューリング(CFS/EEVDF), 仮想記憶(MMU), ページキャッシュ管理", "全リソースの制御権", "システムコール (fork, execve, mmap, kill)", "Kernel Panic (全停止)", "cgroups v2標準対応、Page Cacheの動的割当"),
        ("OSシステムデーモン", "systemd", "PID 1 (init)", "システム初期化、cgroups管理、サービスデーモンの死活監視", "ごく僅かなCPU/RAM", "シグナル (SIGCHLD), D-Bus", "OS停止", "RHEL 9のコアマネージャ"),
        ("OSシステムデーモン", "amazon-ssm-agent", "PID 850 (Daemon)", "AWS Systems Managerとの通信、リモート管理・インベントリ収集", "低 (定期的に微小バースト)", "HTTPS (AWS API)", "リモート管理不可", "性能試験中のCPUノイズ要因になり得る"),
        ("OSシステムデーモン", "amazon-cloudwatch-agent", "PID 890 (Daemon)", "OSメトリクス・ログの収集およびCloudWatchへの定期送信", "低 (ログ量に応じメモリ増)", "HTTPS, Local Pipe", "監視メトリクス欠損", "ログ流量激増時のディスクI/O競合に注意"),
        ("OSシステムデーモン", "rsyslog / systemd-journald", "PID 910, 920", "OS・システムログの集約および永続化", "低～中 (ログ出力頻度依存)", "UNIX Domain Socket (/dev/log)", "ログ欠落", "子JVMが標準出力を大量に出す場合のボトルネック"),
        ("OSシステムデーモン", "chronyd", "PID 780 (Daemon)", "高精度時刻同期デーモン (NTP / Amazon Time Sync)", "ごく僅か", "UDP 123", "時刻ズレによる不整合", "分散バッチにおけるタイムスタンプ整合性に必須"),
        ("親JVM (司令塔)", "OpenJDK 21 (Parent Master)", "PID 1001 (User Space)", "バッチキュー監視、子JVMの生成・並列度制御(最大10)、終了監視、戻り値集計", "低Heap (1～2GB), 低CPU", "ProcessBuilder, 標準入出力パイプ, DB状態管理", "バッチ全体の進行停止 (子JVMは起動中のみ継続)", "自身では重い業務ロジックを一切実行しないのが鉄則"),
        ("子JVM #01 (Worker 1)", "OpenJDK 21 (Child Batch)", "PID 1010 (User Space)", "実際の業務バッチ処理① (DB読込、計算、帳票作成、外部API呼出等)", "独立Heap (-Xmx) + Native", "引数(ARG), 環境変数, DB, 終了コード(exit)", "当該ジョブ①のみ異常終了 (他9並列・親は無傷)", "独立した仮想記憶空間(MMU)で完全分離"),
        ("子JVM #02 (Worker 2)", "OpenJDK 21 (Child Batch)", "PID 1011 (User Space)", "実際の業務バッチ処理②", "独立Heap (-Xmx) + Native", "同上", "当該ジョブ②のみ異常終了", "同上"),
        ("子JVM #03～#09", "OpenJDK 21 (Child Batches)", "PID 1012～1018", "各業務バッチ処理③～⑨ (並列稼働)", "各独立Heap + Native", "同上", "個別ジョブのみ異常終了", "各JVMでGCスレッド・JITコンパイラが自律稼働"),
        ("子JVM #10 (Worker 10)", "OpenJDK 21 (Child Batch)", "PID 1019 (User Space)", "実際の業務バッチ処理⑩", "独立Heap (-Xmx) + Native", "同上", "当該ジョブ⑩のみ異常終了", "全10並列が同時に高負荷になるとvCPU/メモリ競合発生"),
        ("OSカーネルスレッド", "kswapd0 / kcompactd0", "PID 50, 51 (Kernel Thread)", "空きメモリ枯渇時のページ回収(Reclaim)、メモリ断片化解消", "カーネルCPU", "内部カーネルイベント", "システム全体の激しいスローダウン", "メモリ枯渇時にCPUを急激に消費してGCと二重苦に")
    ]

    for r_idx, r_data in enumerate(arch_details, 4):
        fill_cur = fill_light_gray if "親JVM" in r_data[1] else (fill_zebra_even if r_idx % 2 == 0 else fill_zebra_odd)
        for c_idx, val in enumerate(r_data, 1):
            cell = ws3.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_body_bold if c_idx in (1, 2) or "親JVM" in r_data[1] else font_body
            cell.fill = fill_cur
            cell.border = border_cell
            cell.alignment = align_center if c_idx in (1, 3) else align_left
        ws3.row_dimensions[r_idx].height = 24

    auto_fit_columns(ws3, max_cols=8)

    # =============================================================
    # Sheet 4: プロセスメモリ・vCPU構造 (Process, Memory & vCPU)
    # =============================================================
    ws4 = wb.create_sheet(title="プロセスメモリ・vCPU構造")
    ws4.views.sheetView[0].showGridLines = True

    ws4.merge_cells("A1:H1")
    ws4["A1"] = "Linux OS × OpenJDK 21 プロセス・メモリ・vCPU 内部構造詳細"
    ws4["A1"].font = font_title
    ws4.row_dimensions[1].height = 30

    ws4.merge_cells("A3:H3")
    ws4["A3"] = "1. プロセス・スレッド構造 (OS Native Thread と OpenJDK 21 の関係性)"
    ws4["A3"].font = font_sec_hdr
    ws4["A3"].fill = fill_slate_med
    ws4["A3"].alignment = align_left
    ws4.row_dimensions[3].height = 24

    headers4_1 = ["項目", "OS視点 (Linux RHEL 9.8)", "Java視点 (OpenJDK 21)", "11プロセス時の実態", "性能上の勘所・リスク", "調査コマンド"]
    style_header_row(ws4, 4, headers4_1, fill=fill_charcoal_dark)
    ws4.merge_cells("F4:H4")

    proc_data = [
        ("プロセス単位", "task_struct, 独立した仮想アドレス空間(mm_struct), PID", "Java仮想マシン(JVM)インスタンス 1台", "合計11個の独立プロセス (親1 + 子10)", "MMUによるハードウェアレベルのメモリ完全保護", "ps -ef | grep java"),
        ("スレッドモデル", "Native POSIXスレッド (NPTL)。clone()システムコールで生成", "java.lang.Thread (プラットフォームスレッド)", "1 JVMあたり30～50本。11 JVMで総計300～500本以上のOSスレッド", "スレッド増によるOSスケジューラ負荷・スタック消費", "ps -eLF | grep java | wc -l"),
        ("GCスレッド", "OSスレッドとしてCPU時間を消費 (SCHED_OTHER)", "ParallelGCThreads, ConcGCThreads", "デフォルトのままだと 10 JVM × N本 でGCスレッドが大爆破", "複数JVMが同時にGCに入った際の激しいCPU競合", "top -H -p <PID>"),
        ("JITコンパイラ", "C1 (Client) / C2 (Server) コンパイラスレッド", "バックグラウンドでJavaバイトコードをネイティブ機械語に翻訳", "子JVM起動直後に全プロセスでJITが一斉稼働", "起動直後数秒～十数秒のCPUスパイク(コンパイルストーム)", "jcmd <PID> Compiler.queue"),
        ("Virtual Threads", "OSは関与しない (少数のキャリアスレッド上で多重化)", "Java 21新機能 (Project Loom)。数万本の軽量スレッド稼働可能", "子JVM内のI/O待ち処理に適用可能", "ブロッキングI/O時のOSスレッド消費を劇的に削減", "jcmd <PID> Thread.dump_to_file")
    ]

    for r_idx, r_val in enumerate(proc_data, 5):
        fill_cur = fill_zebra_even if r_idx % 2 == 0 else fill_zebra_odd
        for c_idx, val in enumerate(r_val[:5], 1):
            cell = ws4.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_body_bold if c_idx == 1 else font_body
            cell.fill = fill_cur
            cell.border = border_cell
            cell.alignment = align_center if c_idx == 1 else align_left
        
        ws4.merge_cells(start_row=r_idx, start_column=6, end_row=r_idx, end_column=8)
        c6 = ws4.cell(row=r_idx, column=6, value=r_val[5])
        c6.font = font_body_code
        c6.alignment = align_left
        for c in range(6, 9):
            ws4.cell(row=r_idx, column=c).border = border_cell
            ws4.cell(row=r_idx, column=c).fill = fill_cur
        ws4.row_dimensions[r_idx].height = 26

    sec2_start = 12
    ws4.merge_cells(f"A{sec2_start}:H{sec2_start}")
    ws4[f"A{sec2_start}"] = "2. JVMメモリ内部内訳 (なぜ「子JVM 10並列のメモリ ＝ 10 × -Xmx」ではないのか？)"
    ws4[f"A{sec2_start}"].font = font_sec_hdr
    ws4[f"A{sec2_start}"].fill = fill_slate_med
    ws4[f"A{sec2_start}"].alignment = align_left
    ws4.row_dimensions[sec2_start].height = 24

    headers4_2 = ["メモリ領域", "分類", "制御パラメータ", "消費メモリの目安", "10並列時の影響", "枯渇時の挙動 / 注意点", "調査コマンド / 確認方法"]
    style_header_row(ws4, sec2_start + 1, headers4_2, fill=fill_charcoal_dark)
    ws4.merge_cells(f"G{sec2_start + 1}:H{sec2_start + 1}")

    mem_data = [
        ("Java Heap (ヒープ)", "Heap領域", "-Xms, -Xmx", "設定値依存 (例: 2GB～8GB)", "10 × 設定値 (例: 10 × 4GB = 40GB)", "OutOfMemoryError: Java heap space。子プロセスのみ停止", "jstat -gcutil <PID> / jcmd GC.heap_info"),
        ("Metaspace (メタスペース)", "Non-Heap (Native)", "-XX:MetaspaceSize\n-XX:MaxMetaspaceSize", "150MB～300MB / JVM", "10 × 250MB ≒ 2.5GB (重複消費)", "未設定だと無制限にNativeメモリを食い尽くすリスク", "jcmd <PID> VM.metaspace"),
        ("CodeCache (コードキャッシュ)", "Non-Heap (Native)", "-XX:ReservedCodeCacheSize", "64MB～240MB / JVM", "10 × 128MB ≒ 1.3GB (重複消費)", "枯渇するとJITコンパイルが停止し極端に低速化", "jcmd <PID> Compiler.codelist"),
        ("Thread Stacks (スタック)", "Off-Heap (Native)", "-Xss (デフォルト1024KB)", "1MB × スレッド数 (例: 40MB)", "10 × 40MB ≒ 400MB", "スレッド数が多いと直線的にOSメモリを圧迫", "jcmd <PID> Thread.print"),
        ("Direct Memory (NIO/Netty)", "Off-Heap (Native)", "-XX:MaxDirectMemorySize", "バッファ依存 (例: 100MB～500MB)", "10 × 200MB ≒ 2.0GB", "OutOfMemoryError: Direct buffer memory", "jcmd <PID> VM.native_memory detail"),
        ("GC管理構造体 (CardTable等)", "JVM Native", "GC種別依存 (G1GCはHeapの約2～5%)", "Heap 4GB時で約100～200MB", "10 × 150MB ≒ 1.5GB", "G1GCのRemembered Set肥大化によるNative圧迫", "jcmd <PID> VM.native_memory detail"),
        ("glibc / Cランタイム", "Native", "MALLOC_ARENA_MAX等", "アリーナ毎に数十MB～数百MB", "10 × 100MB ≒ 1.0GB", "glibcのマルチスレッドメモリ断片化問題", "pmap -x <PID>"),
        ("【合計】1プロセス実メモリ(RSS)", "合計", "上記全領域の合算", "Heap 4GB設定時 ≒ 実効 4.8GB～5.2GB", "10並列で 48GB～52GB (Heap単体比で+8～12GB超！)", "OS物理RAMを超えるとLinux OOM Killer発動", "ps -o pid,vsz,rss,comm -p <PID>")
    ]

    for r_idx, r_val in enumerate(mem_data, sec2_start + 2):
        fill_cur = fill_calc_total if "【合計】" in r_val[0] else (fill_zebra_even if r_idx % 2 == 0 else fill_zebra_odd)
        for c_idx, val in enumerate(r_val[:6], 1):
            cell = ws4.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_body_bold if c_idx == 1 or "【合計】" in r_val[0] else font_body
            cell.fill = fill_cur
            cell.border = border_cell
            cell.alignment = align_center if c_idx in (1, 2) else align_left
        
        ws4.merge_cells(start_row=r_idx, start_column=7, end_row=r_idx, end_column=8)
        c7 = ws4.cell(row=r_idx, column=7, value=r_val[6])
        c7.font = font_body_code
        c7.alignment = align_left
        for c in range(7, 9):
            ws4.cell(row=r_idx, column=c).border = border_cell
            ws4.cell(row=r_idx, column=c).fill = fill_cur
        ws4.row_dimensions[r_idx].height = 26

    sec3_start = sec2_start + len(mem_data) + 3
    ws4.merge_cells(f"A{sec3_start}:H{sec3_start}")
    ws4[f"A{sec3_start}"] = "3. AWS EC2 vCPU と Linux スケジューリング (CFS / EEVDF) の動作原理"
    ws4[f"A{sec3_start}"].font = font_sec_hdr
    ws4[f"A{sec3_start}"].fill = fill_slate_med
    ws4[f"A{sec3_start}"].alignment = align_left
    ws4.row_dimensions[sec3_start].height = 24

    headers4_3 = ["概念", "ハードウェア / カーネル層の動作原理", "10子JVM並列環境における挙動とリスク", "最適化指針 / チューニング策"]
    style_header_row(ws4, sec3_start + 1, headers4_3, fill=fill_charcoal_dark)
    ws4.merge_cells(f"D{sec3_start + 1}:H{sec3_start + 1}")

    vcpu_details = [
        ("EC2 vCPU の正体",
         "1 vCPU ＝ 物理CPUコアの1ハードウェアスレッド (Intel HT / AMD SMT)。物理演算器(ALU/FPU)やL1/L2キャッシュをペアのスレッドと共有している。",
         "vCPUが100%近くまで回り切ると、キャッシュライン競合やパイプライン失速が発生し、1コアあたりの実効スループットが低下する。",
         "重い計算処理（CPUバウンド）がメインのバッチでは、並列数10に対して最低でも 16 vCPU～32 vCPU インスタンス（c6i/m6i.4xlarge等）を選定する。"),
        ("Linux スケジューラ (CFS/EEVDF)",
         "完全公平スケジューラ(CFS)は、各スレッドの仮想実行時間(vruntime)を赤黒木で管理し、最も実行時間の短いスレッドにCPUタイムスライス(数ms)を公平に割り当てる。",
         "11 JVMから合計400スレッドが競合すると、タイムスライスの細切れ化が発生。実行可能待ち行列(runq-sz)が急激に伸長する。",
         "sar -q の runq-sz を監視。runq-sz ＞ vCPU数 の状態が恒常化している場合は、並列数過多またはvCPU不足の確証となる。"),
        ("コンテキストスイッチ (cswch)",
         "CPUが別のスレッドに実行権を切り替える際、レジスタの退避・復元、TLBフラッシュ、L1/L2キャッシュの無効化（キャッシュスラッシング）が発生する。",
         "自発的切替(Voluntary: I/O待ち)と非自発的切替(Involuntary: 割当時間切れによる強制割込)がある。後者が激増するとCPU時間の多くが切替オーバーヘッドで浪費される。",
         "sar -w の cswch/s を監視。毎秒5万～10万回を超える激しいスパイクが見られる場合、子JVMのGCスレッド数(-XX:ParallelGCThreads=2等)を明示制限する。"),
        ("AWS Nitro によるCPU保護",
         "Nitroカードがネットワーク・ストレージ暗号化・管理タスクを専用ハードウェアに完全オフロードするため、ホストOS/ハイパーバイザによるCPUスティール(%steal)は原則ゼロ。",
         "もし sar -u で %steal ＞ 0 が継続して観測される場合は、T系インスタンスのバーストクレジット枯渇、またはAWS基盤の極度の混雑(Noisy Neighbor)を示す。",
         "本番バッチ環境にはT系（バースト可能）ではなく、C系/M系/R系などの固定性能インスタンスを採用することが必須条件。")
    ]

    for r_idx, r_val in enumerate(vcpu_details, sec3_start + 2):
        fill_cur = fill_zebra_even if r_idx % 2 == 0 else fill_zebra_odd
        ws4.cell(row=r_idx, column=1, value=r_val[0]).alignment = align_center
        ws4.cell(row=r_idx, column=2, value=r_val[1]).alignment = align_left
        ws4.cell(row=r_idx, column=3, value=r_val[2]).alignment = align_left
        
        ws4.merge_cells(start_row=r_idx, start_column=4, end_row=r_idx, end_column=8)
        c4 = ws4.cell(row=r_idx, column=4, value=r_val[3])
        c4.alignment = align_left
        
        for c in range(1, 9):
            cell = ws4.cell(row=r_idx, column=c)
            cell.font = font_body_bold if c == 1 else font_body
            cell.fill = fill_cur
            cell.border = border_cell
        ws4.row_dimensions[r_idx].height = 44

    auto_fit_columns(ws4, max_cols=8)
    ws4.column_dimensions["A"].width = 24
    ws4.column_dimensions["B"].width = 28
    ws4.column_dimensions["C"].width = 26
    ws4.column_dimensions["D"].width = 30

    # =============================================================
    # Sheet 5: 性能試験コマンド完全詳解 (sar & Investigation Guide)
    # =============================================================
    ws5 = wb.create_sheet(title="性能試験コマンド完全詳解")
    ws5.views.sheetView[0].showGridLines = True

    ws5.merge_cells("A1:H1")
    ws5["A1"] = "性能試験・トラブルシューティング調査コマンド完全マニュアル (sar / pidstat / vmstat)"
    ws5["A1"].font = font_title
    ws5.row_dimensions[1].height = 30

    ws5.merge_cells("A3:H3")
    ws5["A3"] = "1. sar (sysstat) オプション別 監視項目・危険判定基準・バッチ特有の調査ポイント"
    ws5["A3"].font = font_sec_hdr
    ws5["A3"].fill = fill_slate_med
    ws5["A3"].alignment = align_left
    ws5.row_dimensions[3].height = 24

    headers5_1 = ["コマンド", "対象リソース", "主要メトリクス項目", "メトリクスの詳細な意味", "性能試験における正常値の目安", "危険シグナル・閾値判定", "親子JVMバッチ環境特有の診断アクション"]
    style_header_row(ws5, 4, headers5_1, fill=fill_charcoal_dark)
    ws5.merge_cells("G4:H4")

    sar_records = [
        ("sar -u 1", "CPU使用率", "%user\n%system\n%iowait\n%steal\n%idle",
         "・%user: Javaアプリ(JIT機械語含む)の実行時間\n・%system: システムコール・コンテキストスイッチ時間\n・%iowait: ディスクI/O完了待ちのCPU遊休時間\n・%steal: 物理ホストが他VMにCPUを奪われた時間",
         "・%user: 70～85% (高負荷バッチ稼働中)\n・%system: 15%未満\n・%iowait: 5%未満\n・%steal: 0.0%",
         "・%system ＞ 20%: スレッド過多・過剰I/O\n・%iowait ＞ 15%: EBS IOPS/スループット枯渇\n・%steal ＞ 1%: インスタンスタイプ選定ミス\n・%idle ＝ 0%: 完全飽和・処理遅延発生",
         "10個の子JVMのCPU合計と親JVMのCPUを確認。%systemが高い場合は後述の pidstat -w でスイッチ頻度を特定。%iowait高騰時はログ書き込みやDBフェッチの非同期化を検討。"),
        ("sar -q 1", "実行キュー &\n負荷平均", "runq-sz\nplist-sz\nldavg-1\nldavg-5\nldavg-15",
         "・runq-sz: CPU実行待ちのタスク(スレッド)数\n・plist-sz: システム全体の総タスク数\n・ldavg-1/5/15: 直近1/5/15分の負荷平均 (CPU実行待ち ＋ ディスクI/O待ちの平均)",
         "・runq-sz ＜ vCPU数\n・ldavg-1 ≦ vCPU数 × 1.2",
         "・runq-sz ＞ vCPU数 × 2: CPU待ち行列滞留\n・ldavg-1 ＞ vCPU数 × 2: 著しい過負荷\n・plist-sz が急激に右肩上がり: スレッドリーク",
         "子JVM並列数10に対してvCPUが足りているかを最も雄弁に語る指標。runq-szが定常的にvCPU数を超えている場合、並列度10を8に下げるか、EC2のスケールアップが必要。"),
        ("sar -r 1", "メモリ利用率", "kbmemfree\nkbmemused\n%memused\nkbbuffers\nkbcached\nkbcommit\n%commit",
         "・kbmemfree: 純粋な未使用物理メモリ\n・kbcached: OSページキャッシュ (ファイルI/Oバッファ)\n・kbcommit: 全プロセスが宣言した仮想メモリ総量\n・%commit: 物理RAMに対する要求コミット比率",
         "・kbcached が適度に確保されていること\n・%commit ＜ 100% (安全圏)\n・利用可能メモリ(MemAvailable) ＞ 15%",
         "・kbcached が極端に減少 (数100MB以下)\n・%commit ＞ 120%: 実メモリ枯渇リスク大\n・MemAvailable がゼロに近づく: OOM Killer秒読み",
         "【最重要】子JVM 10個が立ち上がったピーク時に %commit が100%を大きく超えていないか確認。kbcachedが圧縮されるとファイルI/Oが劇的に遅くなるため、OS用の余裕が必須。"),
        ("sar -S 1", "スワップ消費", "kbswpused\n%swpused\nkbswpfree",
         "・kbswpused: スワップ領域の使用量(KB)\n・%swpused: スワップ領域の使用率(%)",
         "・%swpused ＝ 0.0% (完全未使用が理想)",
         "・%swpused ＞ 0% (わずかでも使用されたら危険)",
         "JVMのヒープ領域がスワップアウトされると、ガベージコレクション(GC)のマーク処理時にディスクI/Oが多発し、GC停止時間が100倍～1000倍に跳ね上がる。本番ではスワップ無効化を推奨。"),
        ("sar -B 1", "ページング統計", "pgpgin/s\npgpgout/s\nfault/s\nmajflt/s",
         "・fault/s: マイナーページフォルト (メモリ割当)\n・majflt/s: メジャーページフォルト (ディスク読込を伴うメモリ割当)",
         "・majflt/s ≒ 0\n・fault/s はJVM起動時に一時増加しその後安定",
         "・majflt/s ＞ 10/s: ディスクからのページ回収多発\n(性能が崖のように急落する)",
         "子JVM起動時に大量のクラスロードとヒープ確保でfault/sが跳ね上がるのは正常。しかし稼働中にmajflt/sが発生している場合、メモリ不足でスラッシングを起こしている証拠。"),
        ("sar -w 1", "タスク生成 &\nコンテキスト", "proc/s\ncswch/s",
         "・proc/s: 1秒あたりの新規タスク(プロセス)生成数\n・cswch/s: 1秒あたりの総コンテキストスイッチ回数",
         "・proc/s: 子JVM起動時以外は低水準\n・cswch/s: vCPU数×3,000回/s 以下が目安",
         "・cswch/s ＞ 50,000～100,000/s: スレッド過密\n・proc/s が異常高値: 短命プロセスの異常乱発",
         "親JVMが子JVMを立ち上げる瞬間に proc/s が立ち上がる。cswch/sが異常に高い場合は、子JVM内のスレッドプールの過剰設定や、デフォルトのGC並列スレッド数の競合を疑う。"),
        ("sar -d -p 1", "EBSディスクI/O", "tps\nrkB/s, wkB/s\nareq-sz\nawait\n%util",
         "・tps: 1秒あたりのI/O転送要求回数 (IOPS相当)\n・await: I/O要求の平均待機＋処理時間 (ミリ秒)\n・%util: ディスクデバイスのビジー率 (%)",
         "・await ＜ 5.0ms (SSD/EBS gp3/io2基準)\n・%util ＜ 70%",
         "・await ＞ 20ms: EBS帯域・IOPSバースト枯渇\n・%util ≒ 100%: ディスクI/Oボトルネック完全確定",
         "バッチで大量のCSV読込、DB更新ログ、アプリケーションログ出力を行う場合、EBSのベースラインIOPS/スループット上限に抵触していないか確認。awaitが急増するとJVMがI/O待ちで詰まる。"),
        ("sar -n DEV 1", "ネットワークI/O", "rxpck/s, txpck/s\nrxkB/s, txkB/s",
         "・パケット送受信数および転送レート(KB/s)",
         "・EC2インスタンスタイプのネットワーク上限の70%以内",
         "・帯域上限到達、ドロップパケット(rxdrop/txdrop)発生",
         "RDSや外部S3との大量データ通信を行うバッチにおいて、EC2のネットワーク帯域制限やコネクション枯渇が発生していないか確認。")
    ]

    for r_idx, r_val in enumerate(sar_records, 5):
        fill_cur = fill_zebra_even if r_idx % 2 == 0 else fill_zebra_odd
        for c_idx, val in enumerate(r_val[:6], 1):
            cell = ws5.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_body_code if c_idx == 1 else (font_body_bold if c_idx in (2, 3) else font_body)
            cell.fill = fill_cur
            cell.border = border_cell
            cell.alignment = align_center if c_idx in (1, 2) else align_left
        
        ws5.merge_cells(start_row=r_idx, start_column=7, end_row=r_idx, end_column=8)
        c7 = ws5.cell(row=r_idx, column=7, value=r_val[6])
        c7.font = font_body
        c7.alignment = align_left
        for c in range(7, 9):
            ws5.cell(row=r_idx, column=c).border = border_cell
            ws5.cell(row=r_idx, column=c).fill = fill_cur
        ws5.row_dimensions[r_idx].height = 54

    sec2_s5 = len(sar_records) + 6
    ws5.merge_cells(f"A{sec2_s5}:H{sec2_s5}")
    ws5[f"A{sec2_s5}"] = "2. 性能試験でsarと併用必須の調査コマンド群 (pidstat, vmstat, top, jcmd等)"
    ws5[f"A{sec2_s5}"].font = font_sec_hdr
    ws5[f"A{sec2_s5}"].fill = fill_slate_med
    ws5[f"A{sec2_s5}"].alignment = align_left
    ws5.row_dimensions[sec2_s5].height = 24

    headers5_2 = ["コマンド構文", "調査対象レイヤー", "出力内容と目的", "注目すべきカラム", "親子JVMバッチ環境での具体的な活用シーン"]
    style_header_row(ws5, sec2_s5 + 1, headers5_2, fill=fill_charcoal_dark)
    ws5.merge_cells(f"E{sec2_s5 + 1}:H{sec2_s5 + 1}")

    companion_cmds = [
        ("pidstat -u -r -p ALL 1", "プロセス個別\n(CPU・メモリ)",
         "全プロセスのCPU使用率とRSS(実物理メモリ)を1秒毎にリアルタイム出力",
         "%CPU, %MEM, RSS, VSZ, Command",
         "【最重要】10個の子JVMのうち、どのPIDが暴走してCPU/メモリを食っているかを特定する神コマンド。sarはOS全体しか見えないが、pidstatなら犯人プロセスが一目瞭然。"),
        ("pidstat -w -p ALL 1", "プロセス個別\n(コンテキスト)",
         "プロセスごとの自発的・非自発的コンテキストスイッチ頻度を出力",
         "cswch/s (自発的: I/O待ち)\nnvcswch/s (非自発的: CPU剥奪)",
         "nvcswch/s (Non-voluntary) が突出している子JVMは、CPUタイムスライスを強制剥奪されている。vCPU競合の直接的証拠となる。"),
        ("vmstat -w 1", "OS全体パルス\n(リアルタイム)",
         "CPU、メモリ、スワップ、ディスクI/O、コンテキストスイッチの横断サマリ",
         "r (実行待ち行列)\nb (I/Oブロック待ち)\nsi/so (スワップイン/アウト)\nus, sy, id, wa, st",
         "性能試験中のターミナル常時表示に最適。'r'がvCPUを超えているか、'b'がゼロより大きいか(I/O詰まり)、'si/so'が発生していないかを1秒で総合判定。"),
        ("top -H -p <子JVM_PID>", "JVM内部\nスレッド別",
         "特定の子JVMプロセス内の全NativeスレッドのCPU・メモリ消費を一覧表示",
         "PID(スレッドID), %CPU, TIME+",
         "特定のJVMが高負荷な時、どのスレッド（GCスレッドなのか、JITなのか、業務ワーカースレッドなのか）がCPUを消費しているかを突き止める。"),
        ("dmesg -T | grep -i oom", "OSカーネルログ\n(異常検知)",
         "LinuxカーネルのOOM Killer発動履歴とキルされたプロセス情報の確認",
         "Out of memory: Killed process <PID> (java)",
         "子JVMが突然 exit code 137 で消滅した場合、Javaのログには何も残らない。dmesgを見ることで「OSによって強制射殺された」決定的一報を確認可能。"),
        ("jcmd <PID> VM.native_memory detail", "JVM Native\nメモリ追跡",
         "NMT (Native Memory Tracking) によるJVM内部の全メモリ領域の詳細内訳",
         "Java Heap, Metaspace, Thread, GC, Internal, Symbol, Native Memory Tracking",
         "事前オプション `-XX:NativeMemoryTracking=summary` が必要。Heap外(Metaspace、DirectMemory、GC領域)のメモリリークを1バイト単位で追跡・特定する。"),
        ("jstat -gcutil <PID> 1000", "JVM GC動態\n(リアルタイム)",
         "JVM各世代(Eden, Survivor, Old, Metaspace)の使用率とGC回数・所要時間",
         "S0, S1, E, O, M, YGC, YGCT, FGC, FGCT, GCT",
         "10個の子JVMのGC挙動をモニタリング。Old領域(O)が100%に張り付いてFull GC(FGC)を連発している子JVMを瞬時に炙り出す。"),
        ("cat /proc/sys/vm/swappiness", "カーネル設定",
         "Linuxカーネルが物理メモリからスワップへの追い出しを好む度合い (0～100)",
         "swappiness 値 (RHEL 9標準は 60)",
         "標準の60のままだと、まだメモリに余裕があるのにPage Cacheを優先してJVMヒープをスワップアウトすることがある。本番では `1` または `10` への変更を推奨。")
    ]

    for r_idx, r_val in enumerate(companion_cmds, sec2_s5 + 2):
        fill_cur = fill_zebra_even if r_idx % 2 == 0 else fill_zebra_odd
        for c_idx, val in enumerate(r_val[:4], 1):
            cell = ws5.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_body_code if c_idx == 1 else (font_body_bold if c_idx in (2, 3) else font_body)
            cell.fill = fill_cur
            cell.border = border_cell
            cell.alignment = align_center if c_idx in (2, 3) else align_left
        
        ws5.merge_cells(start_row=r_idx, start_column=5, end_row=r_idx, end_column=8)
        c5 = ws5.cell(row=r_idx, column=5, value=r_val[4])
        c5.font = font_body
        c5.alignment = align_left
        for c in range(5, 9):
            ws5.cell(row=r_idx, column=c).border = border_cell
            ws5.cell(row=r_idx, column=c).fill = fill_cur
        ws5.row_dimensions[r_idx].height = 42

    auto_fit_columns(ws5, max_cols=8)
    ws5.column_dimensions["A"].width = 28
    ws5.column_dimensions["B"].width = 18
    ws5.column_dimensions["C"].width = 30
    ws5.column_dimensions["D"].width = 24

    # =============================================================
    # Sheet 6: サイジング計算シミュレータ (Sizing Calculator)
    # =============================================================
    ws6 = wb.create_sheet(title="サイジング計算シミュレータ")
    ws6.views.sheetView[0].showGridLines = True

    ws6.merge_cells("A1:G1")
    ws6["A1"] = "EC2インスタンス リソース・サイジング試算シミュレータ (計算式組込)"
    ws6["A1"].font = font_title
    ws6.row_dimensions[1].height = 30

    ws6.merge_cells("A2:G2")
    ws6["A2"] = "※ B列の「設定値」を変更すると、C列の計算式によって必要物理メモリ量と推奨EC2インスタンスが自動再計算されます。"
    ws6["A2"].font = font_subtitle
    ws6.row_dimensions[2].height = 18

    # Table 1: Parameters Input
    ws6.merge_cells("A4:G4")
    ws6["A4"] = "【入力パラメータ】JVMおよびシステム前提条件"
    ws6["A4"].font = font_sec_hdr
    ws6["A4"].fill = fill_slate_med
    ws6["A4"].alignment = align_left
    ws6.row_dimensions[4].height = 24

    param_headers = ["パラメータ項目", "設定値", "単位", "JVM引数 / 根拠", "解説・サイジングの観点"]
    style_header_row(ws6, 5, param_headers, fill=fill_charcoal_dark)
    ws6.merge_cells("E5:G5")

    params = [
        ("子JVM 並列起動数", 10, "プロセス", "運用パラメータ", "本システムの中心要件。同時に稼働する最大子プロセス数"),
        ("子JVM 最大ヒープサイズ", 4.0, "GB / JVM", "-Xmx4g (-Xms4g推奨)", "業務ロジックのデータ保持に必要な最大ヒープ"),
        ("子JVM 最大メタスペース", 0.25, "GB / JVM", "-XX:MaxMetaspaceSize=256m", "ロードされるクラスメタデータ領域の上限設定"),
        ("子JVM コードキャッシュ", 0.125, "GB / JVM", "-XX:ReservedCodeCacheSize=128m", "JITコンパイル済みネイティブ機械語の保存領域"),
        ("子JVM スレッド数", 40, "本 / JVM", "アプリ設計値", "業務ワーカースレッド ＋ GC/JIT等のJVM内部スレッド総計"),
        ("子JVM 1スレッドスタック", 1.0, "MB / スレッド", "-Xss1024k (標準値)", "メソッド呼出階層・ローカル変数のスタック領域"),
        ("子JVM その他Native/GC率", 0.10, "比率 (対Heap)", "G1GC構造体 + DirectMemory", "Card Table, Remembered Sets, NIOバッファ等の推定量 (10%)"),
        ("親JVM メモリ見積もり", 2.0, "GB", "-Xmx1.5g + Native", "司令塔JVMのトータル消費メモリ (Heap + Non-Heap)"),
        ("OSカーネル & デーモン予約", 2.0, "GB", "RHEL 9 + Agents", "systemd, kernel, ssm-agent, cloudwatch-agent等の引当"),
        ("OSページキャッシュ推奨率", 0.15, "比率 (対物理RAM)", "Linux Page Cache", "ディスクI/O(ログ・ファイル操作)を高速化するためのRAM余白"),
        ("子JVM 1プロセスあたり目標vCPU (CPU主体)", 1.2, "vCPU / 子JVM", "CPUバウンド型想定", "暗号化・圧縮・計算処理等で高負荷になる場合の割当目標"),
        ("子JVM 1プロセスあたり目標vCPU (I/O主体)", 0.6, "vCPU / 子JVM", "I/Oバウンド型想定", "DB問い合わせや外部API呼出で待ち時間が多い場合の割当目標")
    ]

    for p_idx, p_val in enumerate(params, 6):
        fill_cur = fill_zebra_even if p_idx % 2 == 0 else fill_zebra_odd
        ws6.cell(row=p_idx, column=1, value=p_val[0]).alignment = align_left
        
        c2 = ws6.cell(row=p_idx, column=2, value=p_val[1])
        c2.alignment = align_right
        if isinstance(p_val[1], float):
            c2.number_format = "0.00" if p_val[1] < 1 else "0.0"
        elif isinstance(p_val[1], int):
            c2.number_format = "#,##0"
            
        ws6.cell(row=p_idx, column=3, value=p_val[2]).alignment = align_center
        ws6.cell(row=p_idx, column=4, value=p_val[3]).alignment = align_left
        
        ws6.merge_cells(start_row=p_idx, start_column=5, end_row=p_idx, end_column=7)
        c5 = ws6.cell(row=p_idx, column=5, value=p_val[4])
        c5.alignment = align_left
        
        for c in range(1, 8):
            cell = ws6.cell(row=p_idx, column=c)
            cell.font = font_body_bold if c in (1, 2) else font_body
            cell.fill = fill_cur
            cell.border = border_cell
        ws6.row_dimensions[p_idx].height = 22

    # Table 2: Calculated Results
    calc_start = 20
    ws6.merge_cells(f"A{calc_start}:G{calc_start}")
    ws6[f"A{calc_start}"] = "【計算結果】子JVM・親JVM・OS全体のメモリおよびvCPU所要量試算"
    ws6[f"A{calc_start}"].font = font_sec_hdr
    ws6[f"A{calc_start}"].fill = fill_slate_med
    ws6[f"A{calc_start}"].alignment = align_left
    ws6.row_dimensions[calc_start].height = 24

    calc_headers = ["試算項目", "計算式 / 参照", "計算値 (Excel数式)", "単位", "判定・安全マージン考察"]
    style_header_row(ws6, calc_start + 1, calc_headers, fill=fill_charcoal_dark)
    ws6.merge_cells(f"E{calc_start + 1}:G{calc_start + 1}")

    calc_rows = [
        ("① 1子JVMあたりスレッドスタック計", "B10 * (B11 / 1024)", "=B10*(B11/1024)", "GB", "スレッド数(40) × スタック(1MB) ＝ 約0.04GB"),
        ("② 1子JVMあたりNative/GC追加分", "B7 * B12", "=B7*B12", "GB", "ヒープ(4GB) × 10% ＝ 約0.40GB"),
        ("③ 1子JVMあたり予想実メモリ (RSS)", "B7 + B8 + B9 + C22 + C23", "=B7+B8+B9+C22+C23", "GB", "ヒープ + メタ + コード + スタック + Native ＝ 約4.81GB/JVM"),
        ("④ 全子JVM(10並列) メモリ小計", "B6 * C24", "=B6*C24", "GB", "10プロセス × 4.81GB ＝ 約48.1GB (ヒープ単体40GB比で+8.1GB増!)"),
        ("⑤ 親JVM ＋ OSデーモン計", "B13 + B14", "=B13+B14", "GB", "親JVM(2GB) ＋ OS引当(2GB) ＝ 4.0GB"),
        ("⑥ アプリ＋システム実メモリ総量", "C25 + C26", "=C25+C26", "GB", "全子JVM ＋ 親JVM ＋ OS ＝ 約52.1GB"),
        ("⑦ ページキャッシュ込み推奨物理RAM", "C27 / (1 - B15)", "=C27/(1-B15)", "GB", "実メモリ所要量をPage Cache(15%)で割り返した必要最小限のRAM"),
        ("★ 必要EC2物理メモリ規格 (目安)", "ROUNDUP(C28, 0)", "=ROUNDUP(C28, 0)", "GB", "この値以上のRAMを搭載するEC2インスタンスを選定必須"),
        ("⑧ 必要vCPU目安 (CPUバウンド型)", "ROUNDUP(B6 * B16 + 2, 0)", "=ROUNDUP(B6*B16+2, 0)", "vCPU", "子10並列×1.2 + 親/OS(2) ＝ 14 vCPU → 16 vCPU推奨"),
        ("⑨ 必要vCPU目安 (I/Oバウンド型)", "ROUNDUP(B6 * B17 + 2, 0)", "=ROUNDUP(B6*B17+2, 0)", "vCPU", "子10並列×0.6 + 親/OS(2) ＝ 8 vCPU → 8～16 vCPU推奨")
    ]

    for i, r_info in enumerate(calc_rows, calc_start + 2):
        row_num = i
        ws6.cell(row=row_num, column=1, value=r_info[0]).alignment = align_left
        ws6.cell(row=row_num, column=2, value=r_info[1]).alignment = align_left
        ws6.cell(row=row_num, column=2).font = font_body_code
        
        c3 = ws6.cell(row=row_num, column=3, value=r_info[2])  # Active Excel formula
        c3.alignment = align_right
        c3.number_format = "0.00" if i not in (calc_start + 9, calc_start + 10, calc_start + 11) else "#,##0"
        
        ws6.cell(row=row_num, column=4, value=r_info[3]).alignment = align_center
        
        ws6.merge_cells(start_row=row_num, start_column=5, end_row=row_num, end_column=7)
        c5 = ws6.cell(row=row_num, column=5, value=r_info[4])
        c5.alignment = align_left
        
        is_highlight = "★" in r_info[0] or "④" in r_info[0] or "⑦" in r_info[0] or "⑧" in r_info[0]
        fill_cur = fill_calc_total if is_highlight else (fill_zebra_even if row_num % 2 == 0 else fill_zebra_odd)
        
        for c in range(1, 8):
            cell = ws6.cell(row=row_num, column=c)
            cell.font = font_body_bold if is_highlight or c == 1 else font_body
            cell.fill = fill_cur
            cell.border = border_total if "★" in r_info[0] else border_cell
        ws6.row_dimensions[row_num].height = 24

    # Table 3: EC2 Instance Type Recommendations
    inst_start = 34
    ws6.merge_cells(f"A{inst_start}:G{inst_start}")
    ws6[f"A{inst_start}"] = "【推奨EC2インスタンスタイプ選定マトリクス】"
    ws6[f"A{inst_start}"].font = font_sec_hdr
    ws6[f"A{inst_start}"].fill = fill_slate_med
    ws6[f"A{inst_start}"].alignment = align_left
    ws6.row_dimensions[inst_start].height = 24

    inst_headers = ["インスタンスタイプ", "vCPU数", "メモリ(GB)", "ネットワーク帯域", "EBS帯域", "本バッチに対する適性判定", "選定理由・推奨ポイント"]
    style_header_row(ws6, inst_start + 1, inst_headers, fill=fill_charcoal_dark)

    inst_choices = [
        ("m6i.4xlarge", "16", "64", "最大 12.5 Gbps", "最大 10,000 Mbps", "◎ 最適 (標準)", "RAM 64GBで試算値(約61.3GB)をジャストカバー。16 vCPUで10並列に十分な演算力を提供。バランス最強。"),
        ("m7i.4xlarge", "16", "64", "最大 12.5 Gbps", "最大 10,000 Mbps", "◎ 最適 (最新世代)", "第4世代Xeon(Sapphire Rapids)。m6i比でIPC向上、OpenJDK 21のベクトル演算やGC性能がさらに向上。"),
        ("r6i.2xlarge", "8", "64", "最大 12.5 Gbps", "最大 10,000 Mbps", "◯ 条件付可 (I/O主体)", "メモリ64GBを確保しつつ低コスト。ただし8 vCPUのため、バッチがI/O待ち中心でCPU負荷が低い場合に限定。"),
        ("r6i.4xlarge", "16", "128", "最大 12.5 Gbps", "最大 10,000 Mbps", "◎ 余裕重視 (大規模)", "RAM 128GBで極めて潤沢なPage Cacheを確保。将来的に子JVM並列数を15～20にスケールアップ可能。"),
        ("c6i.4xlarge", "16", "32", "最大 12.5 Gbps", "最大 10,000 Mbps", "× 不可 (メモリ不足)", "RAM 32GBでは試算値(61GB)に全く足りず、10並列起動時に確実にLinux OOM Killerが発動して死滅する。")
    ]

    for i, inst in enumerate(inst_choices, inst_start + 2):
        fill_cur = fill_calc_total if "◎ 最適" in inst[5] else (fill_zebra_even if i % 2 == 0 else fill_zebra_odd)
        for c_idx, val in enumerate(inst, 1):
            cell = ws6.cell(row=i, column=c_idx, value=val)
            cell.font = font_body_bold if c_idx in (1, 6) else font_body
            cell.fill = fill_cur
            cell.border = border_cell
            cell.alignment = align_center if c_idx in (2, 3, 6) else align_left
        ws6.row_dimensions[i].height = 24

    auto_fit_columns(ws6, max_cols=7)
    ws6.column_dimensions["A"].width = 28
    ws6.column_dimensions["B"].width = 24
    ws6.column_dimensions["C"].width = 18
    ws6.column_dimensions["D"].width = 14
    ws6.column_dimensions["E"].width = 20
    ws6.column_dimensions["F"].width = 20
    ws6.column_dimensions["G"].width = 38

    # =============================================================
    # Sheet 7: 歴史的背景・技術選定理由 (Historical Background)
    # =============================================================
    ws7 = wb.create_sheet(title="歴史的背景・技術選定理由")
    ws7.views.sheetView[0].showGridLines = True

    ws7.merge_cells("A1:G1")
    ws7["A1"] = "Javaバッチ処理アーキテクチャの進化史とマルチプロセス分離の必然性"
    ws7["A1"].font = font_title
    ws7.row_dimensions[1].height = 30

    ws7.merge_cells("A2:G2")
    ws7["A2"] = "メインフレーム時代から、単一JVMスレッド型の挫折、そして最新のクラウド・プロセス分離型に至る技術的必然性の系譜"
    ws7["A2"].font = font_subtitle
    ws7.row_dimensions[2].height = 18

    headers7_hist = ["時代・世代", "主流アーキテクチャ", "代表的技術・製品", "メリット・特徴", "直面した致命的課題・挫折要因", "次世代への進化と教訓"]
    style_header_row(ws7, 4, headers7_hist, fill=fill_charcoal_dark)
    ws7.merge_cells("F4:G4")

    history_records = [
        ("第1世代\n(1970～90年代)", "メインフレーム /\nUNIX C言語・Shell", "IBM JCL, COBOL, C言語, Shellスクリプト",
         "ハードウェア資源を排他的に占有。ジョブごとにOSプロセスを立ち上げ、終了後に資源を完全返却する極めて堅牢なモデル。",
         "高価な専用ハードウェアが必要。開発効率の低さ、移植性の欠如、夜間バッチ時間枠(バッチウィンドウ)の突き抜け。",
         "オープン系への移行とJava言語への期待が高まる。"),
        ("第2世代\n(2000～10年代)", "単一巨大JVM\nマルチスレッド型", "Spring Batch (Single JVM), Quartz, WebSphere / WebLogic Batch",
         "1つの巨大なJavaプロセス内で10～50スレッドを走らせて並列処理。プロセス起動オーバーヘッドがゼロでメモリ効率が良い。",
         "【単一JVMの悲劇】\n① 1つのスレッドのOutOfMemoryErrorでJVM全体がクラッシュ\n② 巨大Heap(32GB超)によるStop-The-World Full GCで全ジョブが数分停止\n③ 静的変数(static)やClassLoaderリークによる長期稼働時のメモリ腐敗\n④ JNIやJDBCドライバのネイティブクラッシュ(SIGSEGV)で全ジョブ即死",
         "「スレッドによる論理分離では、エンタープライズの堅牢性は担保できない」という痛烈な教訓。"),
        ("第3世代\n(2015年代～現在)", "親子JVM\nマルチプロセス分離型\n(本システム)", "親JVM (Orchestrator) ＋ 子JVM (Workers via ProcessBuilder)",
         "【MMUによる完全な物理・仮想メモリ分離】\n① 子JVMがOOMやSIGSEGVで死んでも他9本は無傷で継続\n② 子JVM終了時にOSが全メモリ・ファイル記述子を100%完全回収(リークの根絶)\n③ ジョブごとにJVM引数(-Xmx, GC種別)を個別最適化可能",
         "① プロセス起動時のオーバーヘッド(クラスロード、JITコンパイル)\n② JVM共通領域(Metaspace、CodeCache)の重複によるメモリフットプリント増大\n③ プロセス間の連携がIPC/DB/標準入出力経由となり設計難易度微増",
         "メモリリソースの適切なサイジングと、高速化手法(CDS等)の導入により、現在最も信頼性の高いバッチ基盤として定着。"),
        ("第4世代\n(最新・2023年～)", "クラウドネイティブ\nコンテナ分離 / OpenJDK 21", "AWS ECS on Fargate, Kubernetes BatchJob, OpenJDK 21 (ZGC, CDS)",
         "cgroups v2とコンテナによるリソース厳格制限。OpenJDK 21による起動高速化(AppCDS)、超低遅延Generational ZGC、Virtual ThreadsによるI/O並行化。",
         "コンテナイメージの配布オーバーヘッド、K8s基盤の運用複雑性、コールドスタート時間の増大。",
         "EC2単一インスタンス内での親子JVMマルチプロセス構成は、コンテナオーケストレーションの複雑さを排除しつつ、同一の堅牢性を最もシンプルに享受できる最適解。")
    ]

    for r_idx, r_val in enumerate(history_records, 5):
        fill_cur = fill_light_gray if "第3世代" in r_val[0] else (fill_zebra_even if r_idx % 2 == 0 else fill_zebra_odd)
        for c_idx, val in enumerate(r_val[:5], 1):
            cell = ws7.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_body_bold if c_idx in (1, 2) or "第3世代" in r_val[0] else font_body
            cell.fill = fill_cur
            cell.border = border_cell
            cell.alignment = align_center if c_idx in (1, 2) else align_left
        
        ws7.merge_cells(start_row=r_idx, start_column=6, end_row=r_idx, end_column=7)
        c6 = ws7.cell(row=r_idx, column=6, value=r_val[5])
        c6.font = font_body
        c6.alignment = align_left
        for c in range(6, 8):
            ws7.cell(row=r_idx, column=c).border = border_cell
            ws7.cell(row=r_idx, column=c).fill = fill_cur
        ws7.row_dimensions[r_idx].height = 68

    auto_fit_columns(ws7, max_cols=7)
    ws7.column_dimensions["A"].width = 18
    ws7.column_dimensions["B"].width = 24
    ws7.column_dimensions["C"].width = 28
    ws7.column_dimensions["D"].width = 34
    ws7.column_dimensions["E"].width = 38
    ws7.column_dimensions["F"].width = 32

    # =============================================================
    # Sheet 8: チューニング・推奨設定パラメータ (Tuning Best Practices)
    # =============================================================
    ws8 = wb.create_sheet(title="チューニング・推奨設定パラメータ")
    ws8.views.sheetView[0].showGridLines = True

    ws8.merge_cells("A1:H1")
    ws8["A1"] = "OpenJDK 21 × RHEL 9.8 親子JVMバッチ 推奨設定パラメータ・コマンド集"
    ws8["A1"].font = font_title
    ws8.row_dimensions[1].height = 30

    headers8_tune = ["分類", "パラメータ名 / 設定項目", "推奨設定値", "デフォルト値", "設定目的・性能改善効果", "設定しない場合のリスク"]
    style_header_row(ws8, 3, headers8_tune, fill=fill_charcoal_dark)
    ws8.merge_cells("F3:H3")

    tune_items = [
        ("子JVM起動オプション", "-Xms / -Xmx", "-Xms4g -Xmx4g (同一値)", "動的拡張", "ヒープの初期サイズと最大サイズを一致させ、稼働中のヒープリサイズオーバーヘッドとページフォルトを排除。", "ヒープ拡張時にOSページフォールトが頻発しバッチ処理が一時的に引っかかる。"),
        ("子JVM起動オプション", "-XX:ParallelGCThreads", "-XX:ParallelGCThreads=2", "vCPU数に応じた値 (13等)", "【最重要】子JVM1個あたりの並行GCスレッド数を制限。10並列時にGCスレッドが爆発してCPUを奪い合うのを防ぐ。", "10個のJVMが各々13本GCスレッドを作ると合計130本のスレッドがvCPUを奪い合いcswch/sが爆発する。"),
        ("子JVM起動オプション", "-XX:ConcGCThreads", "-XX:ConcGCThreads=1", "ParallelGCThreads/4", "バックグラウンドのマーキングスレッド数を最小化し、業務処理へのCPU割り当てを最大化。", "バックグラウンドGCが業務処理のvCPUを過度に圧迫する。"),
        ("子JVM起動オプション", "-XX:MaxMetaspaceSize", "-XX:MaxMetaspaceSize=256m", "無制限 (Native上限)", "メタスペースの肥大化に上限を設け、予期せぬNativeメモリリークによるOS RAM枯渇を未然に防止。", "クラスローダ等のリーク時にNativeメモリを食い尽くし、Linux OOM Killerを招く。"),
        ("子JVM起動オプション", "-XX:ReservedCodeCacheSize", "-XX:ReservedCodeCacheSize=128m", "240m", "JITコンパイルコード領域をバッチの規模に合わせて適正化し、10並列時のメモリ重複を抑制。", "無駄に大きな仮想メモリ予約を行い、過剰コミットを引き起こす。"),
        ("子JVM起動オプション", "-Xlog:gc*", "-Xlog:gc*,gc+phases=debug:file=/var/log/jvm/gc-%p.log:time,uptime,pid:filecount=5,filesize=50M", "無効 (標準出力)", "Unified JVM GC Loggingにより、GC停止時間、世代別利用率、フェーズ別所要時間をプロセスPID毎に独立ファイルに完全記録。", "GCが遅延原因かどうかの事後検証が不可能になる。"),
        ("起動高速化 (AppCDS)", "-XX:ArchiveClassesAtExit\n-XX:SharedArchiveFile", "app-cds.jsa を生成して読込", "無効", "共通クラスメタデータを単一の読み取り専用共有アーカイブ(CDS)からmmap展開。クラスロード時間半減＆RAM節約。", "子JVMが起動するたびに毎回全jarファイルをZIP展開・検証してクラスをロードしCPU/ディスクを浪費。"),
        ("Linuxカーネル (sysctl)", "vm.swappiness", "vm.swappiness = 1 (または 10)", "60 (RHEL 9標準)", "OSがまだメモリに余裕がある段階でJVMヒープをスワップへ追い出すのを阻止。Page CacheよりJVMヒープを保護。", "メモリ圧迫時にJVMヒープがスワップアウトされ、GC発生時にStop-The-Worldが数分間固まる。"),
        ("Linuxカーネル (sysctl)", "vm.overcommit_memory", "vm.overcommit_memory = 0 (標準)", "0", "ヒューリスティックなオーバーコミットを許可。過度な要求は事前にENOMEMで弾く。", "1にすると無制限にコミットを許可しOOM Killerが突如暴走。2にするとJVMの巨大仮想予約が起動失敗する。"),
        ("Linuxカーネル (cgroups)", "cgroups v2 メモリ制限", "memory.high / memory.max", "無制限", "systemdスライスまたはcgroups v2を用いて、バッチ全体に利用可能物理RAMの85%を上限設定。", "バッチ暴走時にOSのsshdや監視エージェントまで巻き込んでサーバー応答不能になる。")
    ]

    for r_idx, r_val in enumerate(tune_items, 4):
        fill_cur = fill_calc_total if "最重要" in r_val[4] else (fill_zebra_even if r_idx % 2 == 0 else fill_zebra_odd)
        for c_idx, val in enumerate(r_val[:5], 1):
            cell = ws8.cell(row=r_idx, column=c_idx, value=val)
            cell.font = font_body_code if c_idx in (2, 3) else (font_body_bold if c_idx == 1 else font_body)
            cell.fill = fill_cur
            cell.border = border_cell
            cell.alignment = align_center if c_idx in (1, 3, 4) else align_left
        
        ws8.merge_cells(start_row=r_idx, start_column=6, end_row=r_idx, end_column=8)
        c6 = ws8.cell(row=r_idx, column=6, value=r_val[5])
        c6.font = font_body
        c6.alignment = align_left
        for c in range(6, 9):
            ws8.cell(row=r_idx, column=c).border = border_cell
            ws8.cell(row=r_idx, column=c).fill = fill_cur
        ws8.row_dimensions[r_idx].height = 36

    auto_fit_columns(ws8, max_cols=8)
    ws8.column_dimensions["A"].width = 20
    ws8.column_dimensions["B"].width = 25
    ws8.column_dimensions["C"].width = 25
    ws8.column_dimensions["D"].width = 20
    ws8.column_dimensions["E"].width = 38

    # Final enforcement of Meiryo UI across ALL sheets and cells
    for sheet in wb.worksheets:
        enforce_meiryo_all(sheet)

    output_filename = "JVM_Batch_Performance_Architecture_Design.xlsx"
    wb.save(output_filename)
    
    # Reload and ensure wb._fonts[0].name is Meiryo UI
    wb_final = openpyxl.load_workbook(output_filename)
    wb_final._fonts[0].name = "Meiryo UI"
    wb_final.save(output_filename)
    print(f"Workbook successfully regenerated with 8 sheets at {output_filename}")

if __name__ == "__main__":
    build_workbook()
